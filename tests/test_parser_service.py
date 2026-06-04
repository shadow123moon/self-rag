import os
import tempfile
import unittest
from pathlib import Path
import base64

import pymupdf
from docx import Document as DocxDocument


class ParserServiceTest(unittest.TestCase):
    def test_parse_txt_returns_single_cleaned_text(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "demo.txt"
            file_path.write_text("R e d i s\r\n\r\n\r\n很快", encoding="utf-8")

            parsed = parse_file_content(str(file_path), ".txt")

            self.assertEqual(parsed.page_numbers, None)
            self.assertEqual(parsed.texts, ["Redis\n\n很快"])

    def test_parse_pdf_returns_page_texts_and_page_numbers(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "demo.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "Redis page one")
            pdf.save(file_path)
            pdf.close()

            parsed = parse_file_content(str(file_path), ".pdf")

            self.assertEqual(parsed.page_numbers, [1])
            self.assertEqual(len(parsed.texts), 1)
            self.assertIn("Redis page one", parsed.texts[0])

    def test_parse_pdf_filters_repeated_headers_and_footers(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "headers.pdf"
            pdf = pymupdf.open()
            for page_number in range(1, 4):
                page = pdf.new_page(width=595, height=842)
                page.insert_text((72, 24), "PaiSmart Confidential")
                page.insert_text((72, 150), f"Redis body page {page_number}")
                page.insert_text((72, 820), f"Page {page_number}")
            pdf.save(file_path)
            pdf.close()

            parsed = parse_file_content(str(file_path), ".pdf")
            combined = "\n".join(parsed.texts)

            self.assertEqual(parsed.page_numbers, [1, 2, 3])
            self.assertIn("Redis body page 1", combined)
            self.assertIn("Redis body page 2", combined)
            self.assertIn("Redis body page 3", combined)
            self.assertNotIn("PaiSmart Confidential", combined)
            self.assertNotIn("Page 1", combined)
            self.assertNotIn("Page 2", combined)
            self.assertNotIn("Page 3", combined)

    def test_parse_pdf_skips_empty_pages(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "empty-page.pdf"
            pdf = pymupdf.open()
            pdf.new_page(width=595, height=842)
            page = pdf.new_page(width=595, height=842)
            page.insert_text((72, 150), "Only the second page has content")
            pdf.save(file_path)
            pdf.close()

            parsed = parse_file_content(str(file_path), ".pdf")

            self.assertEqual(parsed.page_numbers, [2])
            self.assertEqual(len(parsed.texts), 1)
            self.assertIn("Only the second page has content", parsed.texts[0])

    def test_parse_pdf_extracts_tables_as_markdown(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "table.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            xs = [50, 150, 250]
            ys = [50, 90, 130]
            for x in xs:
                page.draw_line((x, ys[0]), (x, ys[-1]))
            for y in ys:
                page.draw_line((xs[0], y), (xs[-1], y))
            page.insert_text((65, 72), "Metric")
            page.insert_text((165, 72), "Value")
            page.insert_text((65, 112), "MRR")
            page.insert_text((165, 112), "0.85")
            pdf.save(file_path)
            pdf.close()

            parsed = parse_file_content(str(file_path), ".pdf")
            text = parsed.texts[0]

            self.assertIn("| Metric | Value |", text)
            self.assertIn("| --- | --- |", text)
            self.assertIn("| MRR | 0.85 |", text)

    def test_parse_pdf_uses_ocr_when_page_has_no_text_blocks(self):
        from services.parser_service import PdfParser

        class FakeOcrProvider:
            def __init__(self):
                self.calls = 0

            def extract_text(self, page):
                self.calls += 1
                return "OCR recognized architecture diagram text"

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "scanned.pdf"
            pdf = pymupdf.open()
            pdf.new_page(width=400, height=300)
            pdf.save(file_path)
            pdf.close()

            ocr_provider = FakeOcrProvider()
            parsed = PdfParser(ocr_provider=ocr_provider).parse(str(file_path))

            self.assertEqual(ocr_provider.calls, 1)
            self.assertEqual(parsed.page_numbers, [1])
            self.assertEqual(parsed.texts, ["OCR recognized architecture diagram text"])

    def test_parse_pdf_does_not_use_ocr_when_page_has_text_blocks(self):
        from services.parser_service import PdfParser

        class FakeOcrProvider:
            def __init__(self):
                self.calls = 0

            def extract_text(self, page):
                self.calls += 1
                return "should not appear"

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "text.pdf"
            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            page.insert_text((72, 72), "Normal text PDF")
            pdf.save(file_path)
            pdf.close()

            ocr_provider = FakeOcrProvider()
            parsed = PdfParser(ocr_provider=ocr_provider).parse(str(file_path))

            self.assertEqual(ocr_provider.calls, 0)
            self.assertIn("Normal text PDF", parsed.texts[0])

    def test_parse_pdf_uses_region_ocr_for_images_on_mixed_page(self):
        from services.parser_service import PdfParser

        class FakeOcrProvider:
            def __init__(self):
                self.page_calls = 0
                self.region_rects = []

            def extract_text(self, page):
                self.page_calls += 1
                return "should not use full page ocr"

            def extract_regions_text(self, page, rects):
                self.region_rects = list(rects)
                return ["diagram shows cache invalidation"]

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "mixed-image.pdf"
            image_bytes = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
            )
            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            page.insert_text((72, 72), "Normal text PDF")
            page.insert_image(pymupdf.Rect(80, 120, 220, 220), stream=image_bytes)
            pdf.save(file_path)
            pdf.close()

            ocr_provider = FakeOcrProvider()
            parsed = PdfParser(ocr_provider=ocr_provider).parse(str(file_path))

            self.assertEqual(ocr_provider.page_calls, 0)
            self.assertEqual(len(ocr_provider.region_rects), 1)
            self.assertEqual(parsed.page_numbers, [1])
            self.assertIn("Normal text PDF", parsed.texts[0])
            self.assertIn("【图片OCR】\ndiagram shows cache invalidation", parsed.texts[0])

    def test_parse_pdf_writes_debug_overlay_when_enabled(self):
        from services.parser_service import PdfParser

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            tmp_path = Path(tmp_dir)
            file_path = tmp_path / "debug-overlay.pdf"
            output_dir = tmp_path / "pdf_debug"

            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            page.insert_text((72, 72), "Normal text PDF")
            pdf.save(file_path)
            pdf.close()

            parsed = PdfParser(debug_overlay_dir=output_dir).parse(str(file_path))

            overlay_path = output_dir / "debug-overlay_page_001.png"
            self.assertIn("Normal text PDF", parsed.texts[0])
            self.assertTrue(overlay_path.exists())
            self.assertGreater(overlay_path.stat().st_size, 0)

    def test_pdf_debug_artifacts_keep_image_rects_without_ocr_text(self):
        from services.parser_service import PdfParser

        class EmptyRegionOcrProvider:
            def extract_text(self, page):
                return ""

            def extract_regions_text(self, page, rects):
                return [""] * len(rects)

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "image-without-ocr-text.pdf"
            image_bytes = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
            )
            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            page.insert_image(pymupdf.Rect(80, 120, 220, 220), stream=image_bytes)
            pdf.save(file_path)
            pdf.close()

            parser = PdfParser(ocr_provider=EmptyRegionOcrProvider())
            with pymupdf.open(file_path) as doc:
                blocks, image_rects = parser._extract_text_blocks_with_debug(doc[0])

            self.assertEqual(blocks, [])
            self.assertEqual(len(image_rects), 1)
            self.assertAlmostEqual(image_rects[0].x0, 80.0)
            self.assertAlmostEqual(image_rects[0].y0, 120.0)

    def test_pdf_debug_merges_nearby_image_rects(self):
        from services.parser_service import PdfParser

        parser = PdfParser()

        merged = parser._merge_image_rects(
            [
                pymupdf.Rect(10, 10, 60, 60),
                pymupdf.Rect(66, 12, 120, 62),
                pymupdf.Rect(240, 10, 300, 70),
            ]
        )

        self.assertEqual(len(merged), 2)
        self.assertEqual((merged[0].x0, merged[0].y0, merged[0].x1, merged[0].y1), (10.0, 10.0, 120.0, 62.0))
        self.assertEqual((merged[1].x0, merged[1].y0, merged[1].x1, merged[1].y1), (240.0, 10.0, 300.0, 70.0))

    def test_parse_file_content_writes_pdf_debug_overlay_from_env(self):
        from services.parser_service import parse_file_content

        old_debug_dir = os.environ.get("PDF_PARSER_DEBUG_OVERLAY_DIR")
        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            tmp_path = Path(tmp_dir)
            file_path = tmp_path / "env-debug.pdf"
            output_dir = tmp_path / "env_debug"

            pdf = pymupdf.open()
            page = pdf.new_page(width=400, height=300)
            page.insert_text((72, 72), "Env controlled debug")
            pdf.save(file_path)
            pdf.close()

            try:
                os.environ["PDF_PARSER_DEBUG_OVERLAY_DIR"] = str(output_dir)
                parsed = parse_file_content(str(file_path), ".pdf")
            finally:
                if old_debug_dir is None:
                    os.environ.pop("PDF_PARSER_DEBUG_OVERLAY_DIR", None)
                else:
                    os.environ["PDF_PARSER_DEBUG_OVERLAY_DIR"] = old_debug_dir

            overlay_path = output_dir / "env-debug_page_001.png"
            self.assertIn("Env controlled debug", parsed.texts[0])
            self.assertTrue(overlay_path.exists())
            self.assertGreater(overlay_path.stat().st_size, 0)

    def test_parse_docx_extracts_paragraphs_and_tables(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "demo.docx"
            doc = DocxDocument()
            doc.add_paragraph("Redis 是一个高性能缓存系统。")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "技术"
            table.cell(0, 1).text = "用途"
            table.cell(1, 0).text = "Redis"
            table.cell(1, 1).text = "缓存"
            doc.save(file_path)

            parsed = parse_file_content(str(file_path), ".docx")

            self.assertEqual(parsed.page_numbers, None)
            self.assertEqual(len(parsed.texts), 1)
            self.assertIn("Redis 是一个高性能缓存系统。", parsed.texts[0])
            self.assertIn("| 技术 | 用途 |", parsed.texts[0])
            self.assertIn("| --- | --- |", parsed.texts[0])
            self.assertIn("| Redis | 缓存 |", parsed.texts[0])

    def test_parse_docx_preserves_paragraph_and_table_order(self):
        from services.parser_service import parse_file_content

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as tmp_dir:
            file_path = Path(tmp_dir) / "ordered.docx"
            doc = DocxDocument()
            doc.add_paragraph("表格前说明")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "结果"
            doc.add_paragraph("表格后总结")
            doc.save(file_path)

            parsed = parse_file_content(str(file_path), ".docx")
            text = parsed.texts[0]

            self.assertLess(text.index("表格前说明"), text.index("| 指标 | 结果 |"))
            self.assertLess(text.index("| 指标 | 结果 |"), text.index("表格后总结"))

    def test_parse_file_content_rejects_unsupported_extension(self):
        from services.parser_service import parse_file_content

        with self.assertRaises(ValueError):
            parse_file_content("demo.xlsx", ".xlsx")


if __name__ == "__main__":
    unittest.main()
